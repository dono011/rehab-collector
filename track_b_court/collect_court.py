#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
트랙 B — 법원 공고 회생업체 수집기

법원 공고 목록에서 회생·파산·매각 공고를 수집해 엑셀로 저장한다.
트랙 A(DART)가 못 잡는 **비상장 중소기업**이 대상이다.

⚠️ 실행 전에 probe.py 를 먼저 돌려 대상 사이트 구조를 확인할 것.
   probe 결과에 맞춰 아래 SITES 설정을 조정한다.

사용법:
    python3 collect_court.py --check      # 사이트 접근 가능 여부만
    python3 collect_court.py              # 수집 실행
    python3 collect_court.py --pages 5    # 5페이지까지
    python3 collect_court.py --all        # 이미 본 건도 포함
    python3 collect_court.py --html 파일   # 저장된 HTML로 오프라인 테스트

법적 준수:
    · robots.txt 를 확인하고 금지 시 중단한다
    · 요청 간 1초 이상 간격을 둔다
    · 공개된 공고만 수집한다
"""

import argparse
import json
import os
import re
import sys
import time
from datetime import datetime
from pathlib import Path
from urllib.parse import urljoin, urlparse
from urllib.robotparser import RobotFileParser

try:
    import requests
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError as e:
    print(f"필요한 라이브러리가 없습니다: {e.name}")
    print("설치:  pip install requests openpyxl beautifulsoup4 --break-system-packages")
    sys.exit(1)

import parsers

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "common"))
import ai_judge

OUT_DIR = HERE / "output"
SEEN_FILE = OUT_DIR / "seen.json"

UA = ("Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/126.0 Safari/537.36")

DELAY = 1.2          # 요청 간격(초) — 서버 부담 최소화
TIMEOUT = 25

# ─────────────────────────────────────────────────────────
# 대상 사이트
#
# 2026-08-20 VPS 정찰로 확인한 실제 주소·파라미터.
#   · robots.txt 기계 판정(urllib.robotparser) 결과 모두 "허용"
#   · GET 파라미터 필터 동작 확인 (business=08 로 검증)
#   · 서버 렌더링 HTML → Selenium 불필요
#
# 표 구조: 번호 | 관할법원 | 업종 | 회사 | 작성일
# ─────────────────────────────────────────────────────────
MA_NOTICE_URL = "https://www.scourt.go.kr/portal/notice/mainfo/MaNoticeList.work"

SITES = [
    {
        "name": "대법원 회생회사 M&A 매각공고",
        "url": MA_NOTICE_URL,
        "parser": "ma_notice",
        "page_param": "pageIndex",
        "params": {"pageSize": "10"},
        "enabled": True,
    },
]

# 업종을 좁혀 받고 싶을 때 쓰는 코드 (--business 08)
#
# ⚠️ 기본값은 "전체"다. 실제 목록을 보면 법원의 업종 분류가 느슨해서
#    전기·플랜트 업체도 상당수가 "기타"로 등록되어 있다.
#    (예: 세아STX엔테크 → 기타, ㈜이노피아테크 → 기타)
#    좁혀 받으면 놓치므로, 전체를 받아 회사명으로 한 번 더 거르는 편이 낫다.
DEFAULT_BUSINESS = None

COLUMNS = [
    "수집일자", "회사명", "사업자번호", "대표자명", "관할법원", "사건번호",
    "신청일", "결정일", "사건상태", "면허/업종", "자본금", "채무액",
    "소재지", "전화번호", "데이터출처", "비고",
]


def log(msg=""):
    print(msg, flush=True)


def load_json(path, default):
    if path.exists():
        try:
            return json.loads(path.read_text(encoding="utf-8"))
        except (ValueError, OSError):
            pass
    return default


def save_json(path, obj):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding="utf-8")


# ─────────────────────────────────────────────────────────
# 수집
# ─────────────────────────────────────────────────────────

def robots_allows(url):
    """robots.txt 확인. (허용여부, 사유) 반환."""
    parsed = urlparse(url)
    robots_url = f"{parsed.scheme}://{parsed.netloc}/robots.txt"
    rp = RobotFileParser()
    try:
        resp = requests.get(robots_url, headers={"User-Agent": UA}, timeout=15)
        if resp.status_code != 200:
            return True, "robots.txt 없음 (제한 없음으로 간주)"
        rp.parse(resp.text.splitlines())
    except requests.RequestException as e:
        return True, f"robots.txt 확인 실패, 진행 ({e.__class__.__name__})"

    allowed = rp.can_fetch(UA, url)
    return allowed, ("허용됨" if allowed else "🔴 robots.txt가 이 경로를 금지함")


def fetch(url, params=None):
    """(html, error) 반환."""
    try:
        r = requests.get(url, params=params, headers={"User-Agent": UA},
                         timeout=TIMEOUT)
    except requests.RequestException as e:
        return None, f"{e.__class__.__name__}: {e}"

    if r.status_code != 200:
        return None, f"HTTP {r.status_code}"

    if r.encoding and r.encoding.lower() != "iso-8859-1":
        return r.text, None
    for enc in ("utf-8", "euc-kr", "cp949"):
        try:
            return r.content.decode(enc), None
        except (UnicodeDecodeError, LookupError):
            continue
    return r.content.decode("utf-8", errors="replace"), None


# ─────────────────────────────────────────────────────────
# 공고 원문 열람
#
# 목록에는 회사명·법원·날짜만 있고 사건번호·관리인·연락처가 없다.
# 전화 확인을 하려면 이 정보가 필요하므로 상세 페이지를 열어 뽑아낸다.
#
# 상세 페이지 구조는 확인되지 않았으므로, 태그를 걷어낸 본문에서
# 정규식으로 찾는 방식을 쓴다. 못 찾으면 원문을 파일로 남겨
# 사람이 직접 읽을 수 있게 한다.
# ─────────────────────────────────────────────────────────

# 다음 항목이 시작되는 지점 — 값을 어디서 끊을지 판단하는 데 쓴다.
_NEXT_FIELD = r"(?=\s*(?:인수|제출|문의|연락|접수|일시|기한|담당|주소|$))"

DETAIL_PATTERNS = [
    ("사건번호", re.compile(r"(20\d{2}\s?(?:간회합|회합|회단|하합|하단|개회)\s?\d+)")),
    # 이름 뒤에 조사가 붙는 경우가 많아(관리인 김철수는), 조사를 경계로 끊는다
    ("관리인", re.compile(r"관리인\s*[:：]?\s*([가-힣]{2,4}?)"
                       r"(?=(?:은|는|이|가|께서|씨)?\s)")),
    ("매각주간사", re.compile(r"매각\s*주간사\s*[:：]?\s*(.{2,25}?)" + _NEXT_FIELD)),
    ("전화", re.compile(r"(0\d{1,2}[-)\s]\d{3,4}[-\s]\d{4})")),
    ("이메일", re.compile(r"([\w.+-]+@[\w-]+\.[\w.]+)")),
    ("제출기한", re.compile(r"(?:제출\s*기한|접수\s*마감|입찰\s*일시)\s*[:：]?\s*"
                        r"(20\d{2}[.\-년]\s?\d{1,2}[.\-월]\s?\d{1,2})")),
]


def fetch_detail(url, save_dir=None, slug=""):
    """공고 원문에서 확인 작업에 필요한 항목을 뽑는다."""
    if not url:
        return {}
    html, err = fetch(url)
    if err:
        return {"_error": err}

    text = re.sub(r"<script.*?</script>|<style.*?</style>", " ", html,
                  flags=re.S | re.I)
    text = re.sub(r"<[^>]+>", " ", text)
    text = re.sub(r"&nbsp;?", " ", text)
    text = " ".join(text.split())

    found = {}
    for name, pattern in DETAIL_PATTERNS:
        m = pattern.search(text)
        if m:
            found[name] = " ".join(m.group(1).split())

    # 정규식이 놓친 내용을 사람이 읽을 수 있도록 원문을 남긴다
    if save_dir and slug:
        safe = re.sub(r"[^\w가-힣]+", "_", slug)[:40].strip("_")
        path = Path(save_dir) / f"{safe}.txt"
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(f"{url}\n\n{text}", encoding="utf-8")
        found["_원문"] = str(path)

    return found


def parse_page(site, html, url):
    """사이트 설정에 맞는 파서를 고른다."""
    if site.get("parser") == "ma_notice":
        return parsers.parse_ma_notice(html, base_url=url)
    return parsers.parse_auto(html, base_url=url)


def crawl_site(site, max_pages, business=None):
    """사이트 1곳에서 레코드를 수집한다."""
    name, url = site["name"], site["url"]
    log(f"\n  ▸ {name}")
    log(f"    {url}")

    allowed, reason = robots_allows(url)
    log(f"    robots.txt: {reason}")
    if not allowed:
        log("    → 이 사이트는 건너뜁니다.")
        return []

    if business:
        log(f"    업종 필터: {business} "
            f"({parsers.BUSINESS_CODES.get(business, '알 수 없는 코드')})")

    records = []
    seen_keys = set()
    for page in range(1, max_pages + 1):
        params = dict(site.get("params") or {})
        if business:
            params["business"] = business
        if site.get("page_param"):
            params[site["page_param"]] = page
        elif page > 1:
            break

        html, err = fetch(url, params)
        if err:
            log(f"    {page}페이지: {err}")
            if page == 1:
                log("    → probe.py 로 실제 주소를 확인하세요.")
            break

        found = parse_page(site, html, url)

        # 같은 내용이 반복되면 마지막 페이지를 지난 것이다
        keys = {(r.get("번호"), r.get("회사명")) for r in found}
        if found and keys <= seen_keys:
            log(f"    {page}페이지: 이전과 동일 → 마지막 페이지")
            break
        seen_keys |= keys

        for r in found:
            r["_source"] = name
        records.extend(found)
        log(f"    {page}페이지: {len(found)}건 (누적 {len(records)})")

        if not found:
            break
        time.sleep(DELAY)

    return records


# ─────────────────────────────────────────────────────────
# 출력
# ─────────────────────────────────────────────────────────

def to_row(rec, today):
    """파서 레코드 → 엑셀 한 줄 (트랙 A와 동일한 16열)."""
    status = rec.get("상태", "")
    date = rec.get("날짜", "")
    score = rec.get("관련도", 0)
    induty = rec.get("업종", "")

    note = [f"[관련도 {score}/3]"]
    if rec.get("번호"):
        note.append(f"공고 {rec['번호']}번")
    note.append(rec.get("제목", ""))
    if rec.get("링크"):
        note.append(f"| 원문: {rec['링크']}")

    return {
        "수집일자": today,
        "회사명": rec.get("회사명", ""),
        "사업자번호": "",                       # 법원 공고에 없음
        "대표자명": "",                         # 법원 공고에 없음
        "관할법원": rec.get("법원", ""),
        "사건번호": rec.get("사건번호", ""),     # 목록에 없으면 상세 확인 필요
        "신청일": date if status == "신청" else "",
        "결정일": date if status not in ("신청", "") else "",
        "사건상태": status,
        "면허/업종": induty or "미분류",         # 법원이 분류한 업종
        "자본금": "",                           # 법원 공고에 없음
        "채무액": "",                           # 법원 공고에 없음
        "소재지": "",                           # 공고 본문 확인 필요
        "전화번호": "",                         # 법원 공고에 없음
        "데이터출처": rec.get("_source", "법원공고"),
        "비고": " ".join(note)[:400],
        "_score": score,
        "_name_score": rec.get("이름점수", 0),
        "_induty_score": rec.get("업종점수", 0),
        "_link": rec.get("링크", ""),
        "_key": f"{rec.get('_source','')}|{rec.get('번호','')}|{rec.get('회사명','')}",
    }


def write_excel(rows, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "법원공고"

    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    hit_fill = PatternFill("solid", fgColor="FFF2CC")
    for row in rows:
        ws.append([row.get(c, "") for c in COLUMNS])
        if row.get("_score", 0) >= 2:
            for cell in ws[ws.max_row]:
                cell.fill = hit_fill

    widths = {"회사명": 26, "비고": 70, "면허/업종": 22,
              "관할법원": 14, "사건번호": 14, "데이터출처": 26}
    for idx, name in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(name, 12)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{ws.max_row}"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# ─────────────────────────────────────────────────────────
# 메인
# ─────────────────────────────────────────────────────────

def run_check():
    log("=" * 62)
    log("  대상 사이트 접근 확인")
    log("=" * 62)
    ok = 0
    for site in SITES:
        log(f"\n  ▸ {site['name']}")
        allowed, reason = robots_allows(site["url"])
        log(f"    robots.txt : {reason}")
        html, err = fetch(site["url"])
        if err:
            log(f"    접근       : ❌ {err}")
            continue
        found = parse_page(site, html, site["url"])
        log(f"    접근       : ✅ HTTP 200 / {len(html):,}자")
        log(f"    파싱       : {len(found)}건 추출")
        if found:
            ok += 1
            for r in found[:5]:
                mark = "★" if r.get("관련도", 0) >= 2 else "·"
                log(f"      {mark} {r['회사명']} | {r.get('법원','')} | "
                    f"{r.get('업종','')} | {r.get('날짜','')}")
            if len(found) > 5:
                log(f"      … 외 {len(found) - 5}건")
        else:
            log("      → 추출 0건. probe.py 로 구조를 확인하세요.")
        time.sleep(DELAY)

    log("\n" + "=" * 62)
    log(f"  사용 가능한 사이트: {ok}/{len(SITES)}")
    log("=" * 62)
    return 0 if ok else 1


def main():
    p = argparse.ArgumentParser(description="법원 공고 회생업체 수집기")
    p.add_argument("--pages", type=int, default=5, help="사이트당 페이지 수 (기본 5)")
    p.add_argument("--all", action="store_true", help="이미 본 건도 포함")
    p.add_argument("--check", action="store_true", help="접근 가능 여부만 확인")
    p.add_argument("--html", help="저장된 HTML 파일로 오프라인 테스트")
    p.add_argument("--business", default=DEFAULT_BUSINESS,
                   help="업종코드로 좁혀 받기 (예: 08 건설·엔지니어링). "
                        "기본은 전체 — 법원 분류가 느슨해 좁히면 놓친다")
    p.add_argument("--codes", action="store_true", help="업종·법원 코드표 출력")
    p.add_argument("--since", help="이 날짜 이후 공고만 (예: 2025-01-01)")
    p.add_argument("--elec", action="store_true",
                   help="전기 관련 후보만 (관련도 2 이상)")
    p.add_argument("--tsv", action="store_true",
                   help="엑셀에 붙여넣을 수 있는 표 형태로 화면 출력")
    p.add_argument("--detail", action="store_true",
                   help="공고 원문을 열어 사건번호·관리인·연락처까지 확보 "
                        "(건수만큼 시간이 걸리므로 --elec 와 함께 쓰세요)")
    p.add_argument("--ai-judge", action="store_true",
                   help="전기 관련 후보(관련도 2 이상)를 AI로 재판정 (관련도/위험신호/판단근거). "
                        "기본 제공자는 Jev — 환경변수 AI_JUDGE_PROVIDER=claude|openai 로 변경 가능")
    args = p.parse_args()

    if args.codes:
        log("업종코드 (--business)")
        for code, name in sorted(parsers.BUSINESS_CODES.items()):
            mark = " ★" if code in parsers.BUSINESS_SCORE else ""
            log(f"  {code}  {name}{mark}")
        log("\n법원코드")
        for code, name in sorted(parsers.COURT_CODES.items(), key=lambda x: x[1]):
            log(f"  {code}  {name}")
        return 0

    today = datetime.now()

    if args.check:
        return run_check()

    # 오프라인 테스트 모드
    if args.html:
        path = Path(args.html)
        if not path.exists():
            log(f"파일이 없습니다: {path}")
            return 1
        html = path.read_text(encoding="utf-8", errors="replace")
        records = parse_page(SITES[0], html, MA_NOTICE_URL)
        log(f"{path.name} → {len(records)}건 추출\n")
        for r in records:
            mark = "★" if r["관련도"] >= 2 else " "
            log(f"  {mark} [{r['상태']:4}] {r['회사명']}  "
                f"{r['법원']} {r['사건번호']} {r['날짜']}")
        return 0

    log("=" * 62)
    log("  트랙 B — 법원 공고 회생업체 수집")
    log("=" * 62)
    log(f"  실행   : {today:%Y-%m-%d %H:%M}")
    log(f"  페이지 : 사이트당 최대 {args.pages}")

    log("\n[1/3] 공고 수집")
    all_records = []
    for site in SITES:
        if not site.get("enabled", True):
            continue
        all_records.extend(crawl_site(site, args.pages, args.business))

    if not all_records:
        log("\n수집된 공고가 없습니다.")
        log("  1) python3 probe.py 로 사이트 구조를 확인하세요")
        log("  2) probe 결과에 맞춰 SITES 의 url 을 수정하세요")
        return 1

    all_records = parsers.dedupe(all_records)
    log(f"\n  → 중복 제거 후 {len(all_records)}건")

    if args.since:
        before = len(all_records)
        all_records = [r for r in all_records if (r.get("날짜") or "") >= args.since]
        log(f"  → {args.since} 이후 {len(all_records)}건 "
            f"({before - len(all_records)}건 제외)")

    if args.elec:
        before = len(all_records)
        all_records = [r for r in all_records if r.get("관련도", 0) >= 2]
        log(f"  → 전기 관련 {len(all_records)}건 ({before - len(all_records)}건 제외)")

    if not all_records:
        log("\n조건에 맞는 공고가 없습니다. --since 날짜를 앞당기거나 --elec 를 빼보세요.")
        return 0

    log("\n[2/3] 신규 선별")
    seen = set() if args.all else set(load_json(SEEN_FILE, []))
    rows = [to_row(r, today.strftime("%Y-%m-%d")) for r in all_records]
    new_rows = [r for r in rows if r["_key"] not in seen]

    if not new_rows:
        log(f"  신규 없음 (기존 {len(rows)}건은 이미 수집됨)")
        log("  전부 다시 보려면 --all")
        return 0
    log(f"  신규 {len(new_rows)}건 (기존 {len(rows) - len(new_rows)}건 제외)")

    new_rows.sort(key=lambda r: -r["_score"])

    if args.detail:
        log(f"\n[2.5/3] 공고 원문 열람 ({len(new_rows)}건)")
        detail_dir = OUT_DIR / "detail"
        for n, row in enumerate(new_rows, 1):
            info = fetch_detail(row.get("_link", ""), detail_dir, row["회사명"])
            if info.get("_error"):
                log(f"  [{n}/{len(new_rows)}] {row['회사명']} — 열람 실패: {info['_error']}")
                time.sleep(DELAY)
                continue

            row["사건번호"] = info.get("사건번호", row.get("사건번호", ""))
            row["_관리인"] = info.get("관리인", "")
            row["_연락처"] = info.get("전화", "") or info.get("이메일", "")
            row["_주간사"] = info.get("매각주간사", "")
            row["_기한"] = info.get("제출기한", "")

            bits = [f"{k} {v}" for k, v in
                    (("사건", row["사건번호"]), ("관리인", row["_관리인"]),
                     ("연락처", row["_연락처"]), ("주간사", row["_주간사"]))
                    if v]
            log(f"  [{n}/{len(new_rows)}] {row['회사명']} — "
                + (" | ".join(bits) if bits else "추출 항목 없음 (원문 저장됨)"))

            extra = [f"관리인 {row['_관리인']}" if row["_관리인"] else "",
                     f"연락처 {row['_연락처']}" if row["_연락처"] else "",
                     f"주간사 {row['_주간사']}" if row["_주간사"] else "",
                     f"제출기한 {row['_기한']}" if row["_기한"] else ""]
            extra = [e for e in extra if e]
            if extra:
                row["비고"] = (" / ".join(extra) + " | " + row["비고"])[:400]
            time.sleep(DELAY)
        log(f"\n  원문 저장 위치: {detail_dir}")

    if args.ai_judge:
        provider = os.environ.get("AI_JUDGE_PROVIDER", "jev")
        targets = [r for r in new_rows if r["_score"] >= 2]
        log(f"\n[AI 판정] provider={provider} — 전기 관련 후보 {len(targets)}건")
        for n, row in enumerate(targets, 1):
            result = ai_judge.judge_candidate(
                row["회사명"], row.get("비고", ""), row.get("_link", ""))
            if not result or result.get("_error"):
                log(f"  [{n}/{len(targets)}] {row['회사명']} — 판정 실패: "
                    f"{(result or {}).get('_error', '응답 없음')}")
                continue
            tag = (f"[AI:{result.get('_provider','')}] "
                   f"관련도={result.get('관련도','')} 위험={result.get('위험신호','')} "
                   f"— {result.get('판단근거','')}")
            row["비고"] = (tag + " | " + row["비고"])[:400]
            log(f"  [{n}/{len(targets)}] {row['회사명']} — {tag}")

    log("\n[3/3] 엑셀 저장")
    out_path = OUT_DIR / f"법원공고_회생업체_{today:%Y%m%d}.xlsx"
    write_excel(new_rows, out_path)
    save_json(SEEN_FILE, sorted(seen | {r["_key"] for r in new_rows}))
    log(f"  → {out_path}")

    hits = [r for r in new_rows if r["_score"] >= 2]
    log("\n" + "=" * 62)
    log(f"  신규 {len(new_rows)}건 저장 | 전기 관련 후보 {len(hits)}건")
    log("=" * 62)

    if args.tsv:
        print_tsv(new_rows)
        return 0

    if hits:
        log("\n  ★ 전기 관련 후보")
        for r in hits:
            src = "이름" if r["_name_score"] >= 2 else "업종"
            log(f"     · {r['회사명']} | {r['관할법원']} | {r['결정일'] or r['신청일']} "
                f"| {r['면허/업종']} ({src})")
        log("\n  ⚠️ 자동 추정입니다. 전기공사업 등록 여부는")
        log("     한국전기공사협회(02-2670-5000) 확인이 필요합니다.")
    log("")
    log("  엑셀로 옮기려면:  같은 명령에 --tsv 를 붙여 실행 후 화면을 복사")
    log("")
    return 0


def print_tsv(rows):
    """
    엑셀에 그대로 붙여넣을 수 있는 탭 구분 표를 화면에 출력한다.

    파일을 서버 밖으로 꺼내기 번거로우므로, 화면을 복사해
    엑셀에 붙이면 칸이 자동으로 나뉘도록 탭으로 구분한다.
    """
    cols = ["회사명", "관할법원", "업종", "공고일", "관련도", "판단근거",
            "사건번호", "관리인", "연락처", "매각주간사", "제출기한",
            "확인여부", "메모", "원문링크"]
    log("")
    log("─" * 62)
    log("  아래 표 전체를 복사해 엑셀 A1 칸에 붙여넣으세요 (칸이 자동으로 나뉩니다)")
    log("─" * 62)
    log("")
    log("\t".join(cols))
    for r in rows:
        basis = ("회사명" if r["_name_score"] >= 2
                 else "업종" if r["_induty_score"] >= 2
                 else "참고")
        log("\t".join([
            r["회사명"],
            r["관할법원"],
            r["면허/업종"],
            r["결정일"] or r["신청일"],
            str(r["_score"]),
            basis,
            r.get("사건번호", ""),
            r.get("_관리인", ""),
            r.get("_연락처", ""),
            r.get("_주간사", ""),
            r.get("_기한", ""),
            "",          # 확인여부 — 직접 채우실 칸
            "",          # 메모 — 직접 채우실 칸
            r.get("_link", ""),
        ]))
    log("")


if __name__ == "__main__":
    try:
        sys.exit(main())
    except KeyboardInterrupt:
        log("\n중단되었습니다.")
        sys.exit(130)
